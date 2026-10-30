"""Run the pinned U-C4N v8 matrix through AgentCAD's public MCP interface.

The upstream task definitions, result contract, runner and scoring are reused.
Upstream validators additionally read finished DXFs; rival code performs no
credited CAD work. The cancellation probe imports the actual server worker in
a separate product-interpreter process because that internal race has no public
MCP staging operation.
Fixture creation and independent file readback belong to the evaluator; all CAD
work being credited belongs to the real server. Partial gates receive zero.
"""

from __future__ import annotations

import argparse
import asyncio
from contextlib import asynccontextmanager
import hashlib
import importlib.metadata
import json
import math
import os
from pathlib import Path
import shutil
import subprocess
import sys
import time

import ezdxf
from mcp import ClientSession, StdioServerParameters
from mcp.client.stdio import stdio_client
from pypdf import PdfReader

PIN = "58a05ddfc34a9832c626e82cd88ecdce8e16a62a"
REPO = Path(__file__).resolve().parents[2]


def git(path: Path, *args: str) -> str:
    return subprocess.check_output(["git", "-C", str(path), *args], text=True).strip()


def digest(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def write_json(path: Path, data) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(data, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def strict_result(task_id, passed, metrics=None, message="", *, partial=False):
    from benchmarks.adapters.base import TaskResult

    return TaskResult(
        task_id,
        "pass" if passed else ("partial" if partial else "fail"),
        100.0 if passed else 0.0,
        message=message,
        metrics=metrics or {},
    )


def grade_roundtrip(before, after, reopened, *, parity_unchanged, bores_moved):
    """Require every original measurement, including independently expected values."""
    for report, distance, failed in [
        (before, math.hypot(45, 25), []),
        (after, math.hypot(48, 25), ["R3"]),
        (reopened, math.hypot(48, 25), ["R3"]),
    ]:
        rows = {r["id"]: r for r in report["constraints"]}
        if set(rows) != {"R1", "R2", "R3"} or report["failed"] != failed:
            return False
        if any(
            abs(rows[key]["actual"] - value) > 1e-6
            for key, value in {"R1": 20, "R2": 90, "R3": distance}.items()
        ):
            return False
        count = report["count"]
        if (
            len(count) != 1
            or count[0]["id"] != "R4"
            or count[0]["status"] != "pass"
            or count[0]["actual"] != 4
        ):
            return False
    return parity_unchanged and bores_moved == 1


async def original_quality(path):
    """Read-only oracle: unchanged upstream all-focus critique/finalize score."""
    from dataclasses import asdict
    from backends.ezdxf_backend import EzdxfBackend
    from engineering import DrawingValidator
    from engineering.critique import run_critique
    from engineering.scoring import combine

    before = digest(path)
    backend = EzdxfBackend()
    await backend.connect()
    try:
        await backend.drawing_open(str(path))
        issues = await run_critique(backend, None)
        validation = await DrawingValidator().run(backend)
        summary = {s: sum(i.severity == s for i in issues) for s in ("error", "warning", "info")}
        result = {
            "issues": [asdict(i) for i in issues],
            "validation": [asdict(i) for i in validation.findings],
            "quality": combine(validation.summary, summary),
        }
    finally:
        await backend.disconnect()
    if digest(path) != before:
        raise RuntimeError("The independent quality oracle changed the drawing")
    result["passed"] = not issues and result["quality"]["score"] >= 90
    return result


class AgentCADAdapter:
    name = "agentcad"

    def __init__(self, python: Path):
        self.python = python
        info = subprocess.check_output(
            [
                str(python),
                "-I",
                "-c",
                "import agentcad,json,sys,importlib.metadata as m; "
                "print(json.dumps({'module':agentcad.__file__,'python':sys.version,"
                "'dependencies':{k:m.version(k) for k in ['mcp','ezdxf','pypdf']}}))",
            ],
            text=True,
        )
        self.server_environment = json.loads(info)
        expected = REPO / "src" / "agentcad" / "__init__.py"
        if Path(self.server_environment["module"]).resolve() != expected.resolve():
            raise RuntimeError(
                "--python must have this worktree installed editable (see scripts/new-worktree.ps1)"
            )
        self.product_sha = git(REPO, "rev-parse", "HEAD")
        self.product_dirty = bool(git(REPO, "status", "--porcelain", "--untracked-files=no"))
        self.source_hashes = {
            str(p.relative_to(REPO)).replace("\\", "/"): digest(p)
            for p in sorted((REPO / "src" / "agentcad").rglob("*.py"))
        }
        self.session = None
        self.current_task = "setup"
        self.calls = 0
        self.completed = []

    def log(self, kind: str, **data):
        with (self.artifact_dir / "events.jsonl").open("a", encoding="utf-8") as stream:
            stream.write(
                json.dumps(
                    {"event": kind, "task": self.current_task, "time": time.time(), **data},
                    ensure_ascii=False,
                )
                + "\n"
            )

    @asynccontextmanager
    async def connect(self, profile="full"):
        env = {
            **os.environ,
            "AGENTCAD_ROOT": str(self.workspace),
            "AGENTCAD_VIEWER": "0",
            "AGENTCAD_LIVE_SYNC": "0",
            "AGENTCAD_ALLOW_SCRIPTS": "0",
            "AGENTCAD_READ_ONLY": "0",
            "AGENTCAD_ALLOWED_OPS": "",
            "AGENTCAD_MAX_ENTITIES": "200000",
            "AGENTCAD_MAX_OPS": "5000",
            "PYTHONUTF8": "1",
            "AGENTCAD_TOOL_PROFILE": profile,
        }
        params = StdioServerParameters(
            command=str(self.python), args=["-I", "-m", "agentcad.mcp_server"], cwd=str(REPO), env=env
        )
        with (self.artifact_dir / "server.stderr.log").open("a", encoding="utf-8") as stderr:
            async with stdio_client(params, errlog=stderr) as (reader, writer):
                async with ClientSession(reader, writer) as session:
                    await session.initialize()
                    yield session

    async def setup(self, artifact_dir):
        self.artifact_dir = artifact_dir
        self.workspace = artifact_dir / "workspace"
        self.workspace.mkdir()
        async with self.connect() as self.session:
            response = await self.session.list_tools()
            self.tools = {
                tool.name: tool.model_dump(by_alias=True, mode="json", exclude_none=True)
                for tool in response.tools
            }
            write_json(artifact_dir / "tools.json", {"tools": list(self.tools.values())})
            self.catalog = await self.call("list_operations", detail=True)
            self.compact = await self.call("list_operations")
            self.library = await self.call("standard_library")
        self.session = None

    async def cleanup(self):
        self.current_task = "finished"

    def metadata(self):
        return {
            "backend": "ezdxf",
            "interface": "MCP stdio (official SDK)",
            "python": str(self.python),
            "product_sha": self.product_sha,
            "product_dirty": self.product_dirty,
            "source_sha256": self.source_hashes,
            "scripts_enabled": False,
            "live_autocad": False,
            "server_environment": self.server_environment,
            "partial_score_policy": "0; only every original gate earns 100",
            "server_lifecycle": "fresh MCP process per task; extra fresh reader where required",
            "mcp_calls": self.calls,
        }

    async def call(self, name, *, allow_error=False, **arguments):
        self.calls += 1
        response = await self.session.call_tool(name, arguments)
        data = response.structuredContent
        if data is None:
            texts = [c.text for c in response.content if c.type == "text"]
            try:
                data = json.loads("\n".join(texts))
            except ValueError:
                data = {"ok": False, "protocol_error": "\n".join(texts)}
        self.log(
            "call", seq=self.calls, tool=name, arguments=arguments, response=data, is_error=response.isError
        )
        if (response.isError or data.get("ok") is False) and not allow_error:
            raise RuntimeError(f"{name}: {data}")
        return data

    async def apply(self, operations, **kwargs):
        return await self.call("apply_operations", path=self.path, operations=operations, **kwargs)

    async def inspect(self, **kwargs):
        return await self.call("inspect_entities", path=self.path, **kwargs)

    async def validate(self, path=None, **kwargs):
        return (await self.call("validate_drawing", path=path or self.path, detail=True, **kwargs))[
            "validation"
        ]

    def read(self, path=None):
        return ezdxf.readfile(self.workspace / (path or self.path))

    async def run_task(self, task):
        from benchmarks.adapters.base import TaskResult

        self.current_task = task.task_id
        self.path = task.task_id + ".dxf"

        def outputs():
            return {
                p.relative_to(self.workspace): digest(p)
                for p in self.workspace.rglob("*")
                if p.is_file() and ".agentcad" not in p.relative_to(self.workspace).parts
            }

        before_files = outputs()
        self.log("task_start")
        print(f"START {task.task_id}", flush=True)
        handler = getattr(self, "task_" + task.task_id, None)
        if handler is None:
            result = TaskResult(
                task.task_id, "unsupported", 0, "No executable adapter playbook for this task"
            )
        else:
            async with self.connect() as self.session:
                if task.task_id not in {"token_budget", "tool_discovery", "cancel_safety"}:
                    await self.call("create_drawing", path=self.path, units="mm")
                result = await handler()
            self.session = None
        # Freeze all new/changed deliverables before another task can reuse a
        # filename. This includes CSV/XLSX/PDF, plant inputs and delivery manifests.
        for relative, sha in outputs().items():
            if before_files.get(relative) == sha:
                continue
            target = self.artifact_dir / "retained" / task.task_id / relative
            target.parent.mkdir(parents=True, exist_ok=True)
            shutil.copy2(self.workspace / relative, target)
            result.artifacts.append(str(target))
        self.completed.append(result.to_dict())
        write_json(self.artifact_dir / "progress.json", self.completed)
        self.log("task_end", result=result.to_dict())
        print(f"END {task.task_id}: {result.status} ({result.score})", flush=True)
        return result

    async def task_core_geometry(self):
        await self.apply(
            [
                {"op": "add_line", "start": [0, 0], "end": [3, 4], "tag": "line"},
                {"op": "add_circle", "center": [10, 10], "radius": 2},
            ]
        )
        line = (await self.inspect(ids=["tag:line"]))["entities"][0]
        msp = self.read().modelspace()
        saved_lines = list(msp.query("LINE"))
        circles = list(msp.query("CIRCLE"))
        passed = (
            line["length"] == 5
            and len(saved_lines) == len(circles) == 1
            and abs((saved_lines[0].dxf.end - saved_lines[0].dxf.start).magnitude - 5) < 1e-6
            and circles[0].dxf.radius == 2
            and tuple(circles[0].dxf.center)[:2] == (10, 10)
        )
        return strict_result("core_geometry", passed, {"line_length": line["length"]})

    async def task_modify_query(self):
        await self.apply([{"op": "add_line", "start": [0, 0], "end": [10, 0], "tag": "line"}])
        await self.apply([{"op": "move", "target": "tag:line", "delta": [5, 2]}])
        line = (await self.inspect(ids=["tag:line"]))["entities"][0]
        saved = self.read().modelspace().query("LINE")[0]
        return strict_result(
            "modify_query",
            line["start"][:2] == [5, 2] and tuple(saved.dxf.start)[:2] == (5, 2),
            {"start": line["start"]},
        )

    async def task_layers_linetypes(self):
        await self.apply(
            [
                {"op": "create_layer", "name": "GEOMETRY"},
                {"op": "create_layer", "name": "CENTERLINE", "linetype": "CENTER"},
            ]
        )
        doc = self.read()
        return strict_result(
            "layers_linetypes",
            "GEOMETRY" in doc.layers and "CENTER" in doc.linetypes,
            {"center_linetype": doc.layers.get("CENTERLINE").dxf.linetype},
        )

    async def task_dimensions(self):
        await self.apply(
            [{"op": "add_linear_dim", "p1": [0, 0], "p2": [50, 0], "base": [25, 10], "layer": "DIM"}]
        )
        dims = self.read().modelspace().query("DIMENSION")
        return strict_result(
            "dimensions",
            len(dims) == 1 and bool(dims[0].dxf.geometry) and abs(dims[0].get_measurement() - 50) < 1e-6,
            {"entity_type": dims[0].dxftype(), "render_block": dims[0].dxf.geometry},
        )

    async def task_table_mleader(self):
        await self.apply(
            [
                {
                    "op": "add_table",
                    "insert": [0, 20],
                    "rows": [["ITEM", "QTY"], ["A", "1"]],
                    "col_widths": [20, 20],
                    "tag": "table",
                },
                {"op": "add_leader", "points": [[0, 0], [10, 10]], "text": "NOTE"},
            ]
        )
        msp = self.read().modelspace()
        texts = {e.dxf.text for e in msp.query("TEXT MTEXT")}
        passed = {"ITEM", "QTY", "A", "1", "NOTE"} <= texts and len(msp.query("LEADER")) == 1
        return strict_result(
            "table_mleader",
            passed,
            {"representations": ["composite grid + TEXT", "LEADER + MTEXT"], "native_table_mleader": False},
            "The pinned reference also accepts composite representations; native TABLE/MLEADER is not claimed.",
        )

    async def task_transactions(self):
        await self.apply([{"op": "add_circle", "center": [0, 0], "radius": 1}])
        await self.call("rollback", path=self.path)
        count = len(self.read().modelspace())
        return strict_result("transactions", count == 0, {"entity_count_after_rollback": count})

    async def task_quality_refiner(self):
        await self.apply(
            [
                {"op": "add_line", "start": [0, 0], "end": [10, 0]},
                {"op": "add_line", "start": [0, 0], "end": [10, 0]},
            ]
        )
        report = await self.validate()
        duplicates = [i for i in report["issues"] if i["code"] == "DUPLICATE_ENTITY"]
        refined = await self.call("refine_drawing", path=self.path)
        return strict_result(
            "quality_refiner",
            bool(duplicates)
            and refined["final_score"] > refined["initial_score"]
            and len(self.read().modelspace()) == 1,
            refined,
        )

    async def task_preflight(self):
        result = await self.call(
            "preflight",
            spec={
                "kind": "plate",
                "name": "Benchmark plate",
                "units": "mm",
                "width": 50,
                "height": 25,
                "tolerance_policy": "ISO 2768-m",
            },
        )
        return strict_result(
            "preflight", result["ready"] and result["spec_hash"].startswith("sha256:"), result
        )

    async def task_dxf_roundtrip(self):
        await self.apply([{"op": "add_line", "start": [0, 0], "end": [10, 10]}])
        output = "roundtrip-export.dxf"
        await self.call("export_drawing", path=self.path, output=output, format="dxf")
        async with self.connect() as fresh:
            response = await fresh.call_tool("read_drawing", {"path": output})
            data = response.structuredContent or json.loads(response.content[0].text)
            self.log("fresh_read", response=data)
        count = len(self.read(output).modelspace())
        return strict_result(
            "dxf_roundtrip", data.get("ok") and count == 1, {"entity_count": count, "fresh_process": True}
        )

    async def task_auditable_delivery(self):
        await self.apply([{"op": "add_line", "start": [0, 0], "end": [10, 0]}])
        result = await self.call(
            "deliver_drawing",
            path=self.path,
            output_dir="delivery",
            formats=["dxf"],
            min_score=0,
            strict=False,
        )
        manifest = json.loads((self.workspace / result["manifest"]).read_text("utf-8"))
        files_ok = all(digest(self.workspace / f["path"]) == f["sha256"] for f in manifest["artifacts"])
        passed = (
            result["status"] == "success"
            and files_ok
            and manifest["source_sha256"] == digest(self.workspace / self.path)
            and len(ezdxf.readfile(self.workspace / manifest["artifacts"][0]["path"]).modelspace()) == 1
        )
        return strict_result("auditable_delivery", passed, result)

    async def task_tool_discovery(self):
        expected = {
            "FILLET": "fillet",
            "BPOLY": "trace_boundary",
            "QSELECT": "inspect_entities",
            "WBLOCK": "block_from_entities",
            "OVERKILL": "remove_duplicates",
            "CHSPACE": "change_space",
        }
        hits = {}
        for query in expected:
            hits[query] = (await self.call("search_tools", query=query, limit=3))["results"]
        return strict_result(
            "tool_discovery", all(expected[q] in [r["name"] for r in hits[q][:3]] for q in hits), hits
        )

    async def task_token_budget(self):
        wire = json.dumps({"tools": list(self.tools.values())}, ensure_ascii=False, separators=(",", ":"))
        full_tokens = max(1, int(round(len(wire) / 4.19)))
        async with self.connect(profile="discovery") as session:
            tools = [
                t.model_dump(by_alias=True, mode="json", exclude_none=True)
                for t in (await session.list_tools()).tools
            ]
            routed = await session.call_tool(
                "call_tool",
                {
                    "name": "preflight",
                    "arguments": {"spec": {"kind": "plate", "units": "mm", "width": 50, "height": 25}},
                },
            )
        write_json(self.artifact_dir / "discovery-tools.json", {"tools": tools})
        wire = json.dumps({"tools": tools}, ensure_ascii=False, separators=(",", ":"))
        tokens = max(1, int(round(len(wire) / 4.19)))
        routed_data = routed.structuredContent or json.loads(routed.content[0].text)
        return strict_result(
            "token_budget",
            tokens < 2000 and tokens < full_tokens and bool(routed_data.get("ready")),
            {
                "discovery_tokens": tokens,
                "full_tokens": full_tokens,
                "catalog_chars": len(wire),
                "advertised_tools": len(tools),
                "ceiling": 2000,
                "tokenizer": "ratio/v1 (estimate)",
                "instructions_included": False,
            },
            "AGENTCAD_TOOL_PROFILE=discovery; default full profile remains compatible. Same upstream ratio estimator.",
        )

    async def task_hatch_islands(self):
        await self.apply(
            [
                {
                    "op": "add_hatch",
                    "points": [[0, 0], [20, 0], [20, 20], [0, 20]],
                    "holes": [[[5, 5], [15, 5], [15, 15], [5, 15]]],
                    "pattern": "ANSI31",
                    "style": "outer",
                }
            ]
        )
        inspected = (await self.inspect(filter={"type": "HATCH"}))["entities"][0]
        hatch = self.read().modelspace().query("HATCH")[0]
        before = (await self.call("measure_entities", path=self.path, target=inspected["id"]))[
            "measurements"
        ][0]
        changed = await self.apply([{"op": "edit_hatch", "target": inspected["id"], "style": "ignore"}])
        after = (await self.call("measure_entities", path=self.path, target=inspected["id"]))["measurements"][
            0
        ]
        passed = (
            len(hatch.paths) == 2
            and hatch.dxf.hatch_style == 1
            and before["area"] == 300
            and after["area"] == 400
            and self.read().entitydb[inspected["id"]].dxf.hatch_style == 2
            and "style" in changed["results"][0]["edits"][0]["changed"]
        )
        return strict_result("hatch_islands", passed, {"before": before, "after": after, "edit": changed})

    async def task_selection_filter(self):
        await self.apply(
            [
                {"op": "add_circle", "center": [5, 5], "radius": 1},
                {"op": "add_circle", "center": [10, 5], "radius": 3},
                {"op": "add_circle", "center": [18, 18], "radius": 0.5},
                {"op": "add_circle", "center": [40, 40], "radius": 1, "layer": "TEXT"},
            ]
        )
        crossing = await self.inspect(window=[[0, 0], [10, 10]])
        by_layer = await self.inspect(filter={"type": "CIRCLE", "layer": "TEXT"})
        window = await self.inspect(window=[[0, 0], [10, 10]], selection_mode="window")
        box = await self.inspect(window=[[0, 0], [20, 20]], selection_mode="window")
        polygon = await self.inspect(polygon=[[0, 0], [20, 0], [0, 20]], selection_mode="window")
        counts = [window["total"], crossing["total"], box["total"], polygon["total"], by_layer["total"]]
        return strict_result(
            "selection_filter",
            counts == [1, 2, 3, 1, 1] and set(by_layer["filtered_by"]) == {"type", "layer"},
            {"counts": counts, "filtered_by": by_layer["filtered_by"]},
        )

    async def task_measure_from_handle(self):
        await self.apply(
            [
                {"op": "add_line", "start": [0, 0], "end": [10, 0]},
                {"op": "add_line", "start": [10, 0], "end": [10, 10]},
                {"op": "add_line", "start": [0, 10], "end": [0, 0]},
                {"op": "add_arc", "center": [5, 10], "radius": 5, "start_angle": 0, "end_angle": 180},
            ]
        )
        await self.apply([{"op": "trace_boundary", "point": [5, 5], "tag": "boundary"}])
        info = (await self.inspect(ids=["tag:boundary"]))["entities"][0]
        measured = (await self.call("measure_entities", path=self.path, target=info["id"]))["measurements"][0]
        area = 100 + math.pi * 25 / 2
        saved = self.read().entitydb[info["id"]]
        vertices = list(saved.get_points("xy"))
        shoelace = (
            abs(sum(a[0] * b[1] - b[0] * a[1] for a, b in zip(vertices, vertices[1:] + vertices[:1]))) / 2
        )
        passed = (
            saved.closed
            and info["closed"]
            and measured["exact"]
            and measured["method"] == "analytic_bulge"
            and abs(measured["area"] - area) < 1e-6
            and shoelace == 100
            and 0.28 < 1 - shoelace / area < 0.29
        )
        return strict_result(
            "measure_from_handle", passed, {"boundary": info, "measurement": measured, "shoelace": shoelace}
        )

    async def task_page_setup_truth(self):
        await self.apply(
            [
                {"op": "set_page_setup", "name": "Layout1", "paper": "ANSI_B"},
                {"op": "add_line", "start": [20, 20], "end": [100, 100], "layout": "Layout1"},
            ]
        )
        await self.call("export_drawing", path=self.path, output="ansi-b.pdf", format="pdf", layout="Layout1")
        page = PdfReader(self.workspace / "ansi-b.pdf").pages[0]
        size = [float(page.mediabox.width) * 25.4 / 72, float(page.mediabox.height) * 25.4 / 72]
        changed = await self.apply(
            [{"op": "set_page_setup", "name": "Layout1", "paper": "A3", "ctb": "monochrome.ctb"}]
        )
        await self.call("export_drawing", path=self.path, output="iso-a3.pdf", format="pdf", layout="Layout1")
        page2 = PdfReader(self.workspace / "iso-a3.pdf").pages[0]
        size2 = [float(page2.mediabox.width) * 25.4 / 72, float(page2.mediabox.height) * 25.4 / 72]
        saved = self.read().layouts.get("Layout1").dxf_layout.dxf
        change = changed["results"][0]
        passed = (
            all(abs(a - b) < 0.01 for a, b in zip(size, [432, 279]))
            and all(abs(a - b) < 0.01 for a, b in zip(size2, [420, 297]))
            and [saved.paper_width, saved.paper_height] == [420, 297]
            and saved.current_style_sheet == "monochrome.ctb"
            and change["ctb_known"]
            and change["page_setup"]["orientation"] == "landscape"
            and change["changed"]["paper_size_mm"] == {"old": [432, 279], "new": [420, 297]}
        )
        return strict_result(
            "page_setup_truth", passed, {"ansi_b_mm": size, "iso_a3_mm": size2, "change": change}
        )

    async def task_pid_roundtrip(self):
        from engineering.pid.spec import EXAMPLE_SPEC

        await self.apply([{"op": "draw_pid", "spec": EXAMPLE_SPEC}])
        result = await self.call("read_engineering", path=self.path, kind="pid")
        return strict_result(
            "pid_roundtrip", grade_pid(result) and len(self.read().modelspace().query("INSERT")) == 5, result
        )

    async def task_pid_file_roundtrip(self):
        from engineering.pid.spec import EXAMPLE_SPEC

        await self.apply([{"op": "draw_pid", "spec": EXAMPLE_SPEC}])
        before = await self.call("read_engineering", path=self.path, kind="pid")
        async with self.connect() as fresh:
            old, self.session = self.session, fresh
            try:
                after = await self.call("read_engineering", path=self.path, kind="pid")
            finally:
                self.session = old
        fields = ["graph", "line_list", "instrument_index", "equipment_list", "critique"]
        same = {name: before[name] == after[name] for name in fields}
        return strict_result(
            "pid_file_roundtrip",
            grade_pid(before) and grade_pid(after) and all(same.values()),
            {"fresh_read": after, "equal": same},
        )

    def plant_fixture(self):
        from tests.fixtures.plant_pair import build_plant_pair

        return build_plant_pair(self.workspace / self.current_task)

    async def task_understand_foreign(self):
        truth = self.plant_fixture()
        layout = await self.call(
            "understand_drawing", path=str(Path(truth["layout"]).relative_to(self.workspace))
        )
        pid = await self.call("understand_drawing", path=str(Path(truth["pid"]).relative_to(self.workspace)))
        read = {r["name"]: r["service"] for r in pid["layers"]}
        passed = (
            layout["units"]["inferred"] == truth["layout_true_unit"]
            and layout["units"]["agree"] is False
            and "INSUNITS" in json.dumps(layout["warnings"])
            and truth["outlier_handle"] in layout["extents"]["outlier_handles"]
            and len(layout["clusters"]) == truth["cluster_count"]
            and all(read.get(name) == value for name, value in truth["layer_services"].items())
        )
        return strict_result("understand_foreign", passed, {"layout": layout, "pid": pid})

    async def task_takeoff_roundtrip(self):
        truth = self.plant_fixture()
        report = await self.call(
            "plant_takeoff",
            pid=str(Path(truth["pid"]).relative_to(self.workspace)),
            layout=str(Path(truth["layout"]).relative_to(self.workspace)),
            services=sorted({r["service"] for r in truth["pipe_runs"]}),
        )
        found = {(r["service"], tuple(sorted(r["tags"]))): r for r in report["pipe"]["runs"]}
        problems = []
        for want in truth["pipe_runs"]:
            got = found.get((want["service"], tuple(want["tags"])))
            if got is None:
                problems.append({"run": want["id"], "reason": "missing"})
                continue
            length = got["layout_m"]
            length_ok = (
                length is None
                if want["layout_length_mm"] is None
                else length is not None and abs(length * 1000 - want["layout_length_mm"]) < 1e-6
            )
            if (
                got["status"] != want["status"]
                or set(got["diameters"]) != set(want["diameters"])
                or not length_ok
            ):
                problems.append({"run": want["id"], "got": got, "want": want})
        loads = {r["tag"]: r for r in report["cable"]["rows"]}
        for want in truth["cable_rows"]:
            got = loads.get(want["tag"])
            if (
                got is None
                or got["panel"] != want["panel"]
                or got["kw"] != want["kw"]
                or got["cable_m"] != want["roundup_m"]
                or got["length_m"] is None
                or abs(got["length_m"] * 1000 - want["length_mm"]) > 1e-6
            ):
                problems.append({"load": want["tag"], "got": got, "want": want})
        passed = (
            not problems
            and len(found) == len(truth["pipe_runs"])
            and len(loads) == len(truth["cable_rows"])
            and report["scale"]["verdict"] == truth["scale_verdict"]
            and report["scale"]["pairs"] == 55
        )
        return strict_result("takeoff_roundtrip", passed, {"report": report, "mismatches": problems})

    async def task_arch_roundtrip(self):
        from engineering.arch.spec import EXAMPLE_SPEC

        await self.apply([{"op": "draw_arch_plan", "spec": EXAMPLE_SPEC}])
        report = await self.call("read_engineering", path=self.path, kind="architecture")
        report["original_quality"] = await original_quality(self.workspace / self.path)
        expected = {"01": 4825 * 5750, "02": 3825 * 5750}
        rooms = {r["number"]: r for r in report["rooms"]}
        doors = sorted(r["tag"] for r in report["schedules"]["doors"])
        windows = sorted(r["tag"] for r in report["schedules"]["windows"])
        passed = (
            set(rooms) == set(expected)
            and len(report["detected_rooms"]) == 2
            and all(
                abs(rooms[n]["area"] - v) < v * 0.001
                and rooms[n]["detected_area"] is not None
                and abs(rooms[n]["detected_area"] - v) < v * 0.001
                for n, v in expected.items()
            )
            and doors == ["D1", "D2"]
            and windows == ["W1", "W2"]
            and len(report["stairs"]) == 1
            and report["critique"] == []
            and report["quality_score"] >= 90
            and report["original_quality"]["passed"]
        )
        return strict_result("arch_roundtrip", passed, report)

    async def task_mech_assembly(self):
        spec = {
            "sheet": {
                "size": "A3",
                "title": "FLANGE COUPLING",
                "drawing_no": "MA-001",
                "drawn_by": "AgentCAD",
                "scale": "1:1",
                "material": "S235JR",
            },
            "parts": [
                {
                    "kind": "revolved",
                    "name": "HUB",
                    "material": "steel",
                    "at": [125, 160],
                    "views": ["front", "section"],
                    "segments": [
                        {"length": 40, "d_outer": 40, "d_inner": 20},
                        {"length": 15, "d_outer": 90, "d_inner": 20},
                    ],
                },
                {
                    "kind": "revolved",
                    "name": "FLANGE",
                    "material": "cast_iron",
                    "at": [345, 160],
                    "views": ["front", "side"],
                    "segments": [{"length": 15, "d_outer": 90, "d_inner": 40}],
                    "features": [
                        {
                            "kind": "hole_circle",
                            "id": "bolts",
                            "x": 7.5,
                            "pcd": 65,
                            "count": 4,
                            "diameter": 13.5,
                            "start_angle": 45,
                        }
                    ],
                },
            ],
            "fasteners": [
                {
                    "part": "hex_bolt" if i % 2 == 0 else "hex_nut",
                    "size": "M12",
                    "insert": [x, 90],
                    "view": "side",
                    **({"length": 60} if i % 2 == 0 else {}),
                }
                for i, x in enumerate([60, 110, 160, 210])
            ],
        }
        await self.apply([{"op": "draw_mechanical_assembly", "spec": spec}])
        report = await self.call("read_engineering", path=self.path, kind="mechanical")
        report["original_quality"] = await original_quality(self.workspace / self.path)
        parts = {p["name"]: p for p in report["parts"]}
        fasteners = {
            r["description"]: r for r in report["bom"] if r["description"] in {"hex_bolt", "hex_nut"}
        }
        holes = [c for c in report["circles"] if c["kind"] == "bolt_hole"]
        doc = self.read()
        passed = (
            report["critique"] == []
            and report["quality_score"] >= 90
            and report["original_quality"]["passed"]
            and set(parts) == {"HUB", "FLANGE"}
            and parts["HUB"]["measured"] == {"length": 55, "diameter": 90}
            and parts["FLANGE"]["measured"] == {"length": 15, "diameter": 90}
            and set(fasteners) == {"hex_bolt", "hex_nut"}
            and all(r["quantity"] == 2 for r in fasteners.values())
            and fasteners["hex_bolt"]["standard"] == "ISO 4014"
            and fasteners["hex_bolt"]["size"] == "M12x60"
            and fasteners["hex_nut"]["standard"] == "ISO 4032"
            and fasteners["hex_nut"]["size"] == "M12"
            and len(holes) == 4
            and all(c["diameter"] == 13.5 for c in holes)
            and len(report["sections"]) == 2
            and len(report["balloons"]) == len(report["bom"])
        )
        passed &= len(doc.modelspace().query("HATCH")) >= 2 and len(doc.modelspace().query("DIMENSION")) >= 3
        return strict_result("mech_assembly", passed, report)

    async def task_requirements_roundtrip(self):
        await self.call(
            "draw_from_spec",
            output=self.path,
            overwrite=True,
            spec={
                "kind": "plate",
                "units": "mm",
                "width": 120,
                "height": 80,
                "material": "steel",
                "thickness": 10,
                "dimensions": "none",
                "center_marks": False,
                "holes": [
                    {"pattern": "single", "tag": tag, "center": [x, y], "diameter": diameter}
                    for tag, x, y, diameter in [
                        ("bore", 60, 40, 20),
                        ("c1", 15, 15, 8),
                        ("c2", 105, 15, 8),
                        ("c3", 105, 65, 8),
                        ("c4", 15, 65, 8),
                    ]
                ],
            },
        )
        constraints = [
            {"id": "R1", "type": "diameter", "a": "tag:bore", "value": 20},
            {"id": "R2", "type": "distance", "a": "tag:c1@center", "b": "tag:c2@center", "value": 90},
            {
                "id": "R3",
                "type": "distance",
                "a": "tag:c1@center",
                "b": "tag:bore@center",
                "value": math.hypot(45, 25),
            },
        ]
        count_req = [{"id": "R4", "type": "hole_count", "diameter": 8, "count": 4}]

        async def measure(session):
            previous = self.session
            self.session = session
            try:
                cr = await self.call(
                    "apply_constraints", path=self.path, constraints=constraints, mode="check"
                )
                vr = await self.validate(requirements=count_req)
                rows = cr["constraints"]
                failed = [r["id"] for r in rows if not r["satisfied"]]
                failed += [r["id"] for r in vr["requirements"] if r["status"] != "pass"]
                return {"failed": sorted(failed), "constraints": rows, "count": vr["requirements"]}
            finally:
                self.session = previous

        before = await measure(self.session)
        doc = self.read()
        from ezdxf import bbox

        parity_before = (len(doc.modelspace()), tuple(bbox.extents(doc.modelspace())))
        bores = [c for c in doc.modelspace().query("CIRCLE") if c.dxf.radius == 10]
        bores[0].dxf.center = (63, 40, 0)
        doc.saveas(self.workspace / self.path)
        saved = self.read()
        parity_after = (len(saved.modelspace()), tuple(bbox.extents(saved.modelspace())))
        after = await measure(self.session)
        async with self.connect() as fresh:
            reopened = await measure(fresh)
        passed = grade_roundtrip(
            before, after, reopened, parity_unchanged=parity_before == parity_after, bores_moved=len(bores)
        )
        return strict_result(
            "requirements_roundtrip",
            passed,
            {
                "document": before,
                "artifact": after,
                "fresh_process": reopened,
                "moved_mm": 3,
                "equivalent_contract": "native check-only constraints R1-R3 + geometry count R4",
            },
        )

    async def task_cancel_safety(self):
        proc = await asyncio.create_subprocess_exec(
            str(self.python),
            "-I",
            str(Path(__file__).with_name("cancel_worker.py")),
            str(self.workspace),
            stdout=asyncio.subprocess.PIPE,
            stderr=asyncio.subprocess.PIPE,
            env={**os.environ, "AGENTCAD_VIEWER": "0", "AGENTCAD_LIVE_SYNC": "0"},
        )
        stdout, stderr = await asyncio.wait_for(proc.communicate(), 30)
        (self.artifact_dir / "cancel-worker.stderr.log").write_bytes(stderr)
        probe = json.loads(stdout)
        read = await self.call("read_drawing", path=self.path)
        probe["entity_count_after"] = read["entity_count"]
        return strict_result(
            "cancel_safety",
            proc.returncode == 0
            and probe["trace"] == ["A:end", "B:run"]
            and probe["overlapped"] is False
            and probe["healthy"]
            and read["entity_count"] == 1,
            probe,
            "Internal held-worker fixture, matching the original's private _async probe; public MCP health read afterwards.",
        )

    async def task_qa_large_drawing(self):
        # Same raw fixture staging as the reference; evaluation only calls the public validator.
        doc = ezdxf.new("R2010", setup=True, units=4)
        doc.layers.add("GEOMETRY")
        doc.layers.add("CONSTRUCTION")
        msp = doc.modelspace()
        construction = {}
        for index in range(1, 10001):
            layer = "CONSTRUCTION" if index == 5001 else "GEOMETRY"
            e = msp.add_line((0, index * 10), (100, index * 10), dxfattribs={"layer": layer})
            if index == 5001:
                construction["at_5001"] = e.dxf.handle
        construction["after_10000"] = msp.add_line(
            (-300, 0), (-200, 0), dxfattribs={"layer": "CONSTRUCTION"}
        ).dxf.handle
        construction["paper"] = (
            doc.layouts.get("Layout1")
            .add_line((0, 0), (50, 0), dxfattribs={"layer": "CONSTRUCTION"})
            .dxf.handle
        )
        a = msp.add_line((500, -50), (600, -50), dxfattribs={"layer": "GEOMETRY"})
        b = msp.add_line((600, -50), (500, -50), dxfattribs={"layer": "GEOMETRY"})
        doc.saveas(self.workspace / self.path)
        report = await self.validate()
        paper = await self.validate(layout="Layout1")
        found = {
            h for report_ in [report, paper] for issue in report_["issues"] for h in issue.get("entities", [])
        }
        missing = [name for name, h in construction.items() if h not in found]
        dup_found = any({a.dxf.handle, b.dxf.handle} <= set(i.get("entities", [])) for i in report["issues"])
        return strict_result(
            "qa_large_drawing",
            not missing and dup_found,
            {
                "construction": construction,
                "missing_construction": missing,
                "duplicate_found": dup_found,
                "issue_codes": sorted({i["code"] for i in report["issues"] + paper["issues"]}),
            },
            "Full model and paper validation called separately; all reference construction findings and the duplicate are required.",
        )

    async def task_export_safety(self):
        from copy import deepcopy
        import csv
        from openpyxl import load_workbook
        from engineering.pid.spec import EXAMPLE_SPEC

        payloads = ["=1+1", "@SUM(A1:A9)", "+cmd|' /C calc'!A0", "-2+3"]
        spec = deepcopy(EXAMPLE_SPEC)
        for line in spec["lines"]:
            if line.get("number"):
                line["number"] = payloads[0]
        await self.apply([{"op": "draw_pid", "spec": spec}])
        await self.call("export_drawing", path=self.path, output="lines.csv", format="csv", table="pid_lines")
        with (self.workspace / "lines.csv").open(encoding="utf-8-sig", newline="") as stream:
            line_rows = list(csv.DictReader(stream))
        pid_ok = [r["line_number"] for r in line_rows].count("'=1+1") == 2
        await self.call("create_drawing", path="export-bom.dxf", units="mm")
        attributes = [{"DESIGNATION": text, "MATERIAL": payloads[3]} for text in payloads[:3]]
        attributes.append(
            {"DESIGNATION": "@NOW()", "STANDARD": "-ISO 4014", "MATERIAL": "=HYPERLINK(1)", "MASS": "-12.5"}
        )
        await self.call(
            "apply_operations",
            path="export-bom.dxf",
            operations=[
                {
                    "op": "define_block",
                    "name": "BENCH_PART",
                    "entities": [{"op": "add_circle", "center": [0, 0], "radius": 1}],
                    "attributes": [
                        {"tag": tag, "insert": [0, i * 3]}
                        for i, tag in enumerate(["DESIGNATION", "STANDARD", "MATERIAL", "MASS"])
                    ],
                },
                *[
                    {"op": "insert_block", "name": "BENCH_PART", "insert": [i * 20, 0], "attributes": a}
                    for i, a in enumerate(attributes)
                ],
            ],
        )
        for fmt in ["csv", "xlsx"]:
            await self.call(
                "export_drawing", path="export-bom.dxf", output=f"export.{fmt}", format=fmt, table="bom"
            )
        with (self.workspace / "export.csv").open(encoding="utf-8-sig", newline="") as stream:
            csv_rows = list(csv.reader(stream))
        wb = load_workbook(self.workspace / "export.xlsx", data_only=False)
        cells = [cell for row in list(wb.active)[1:] for cell in row]
        texts = {c.value for c in cells if c.data_type == "s"}
        numeric_mass = any(c.value == -12.5 and c.data_type == "n" for c in cells)
        formulas = [c.coordinate for c in cells if c.data_type == "f"]
        wb.close()
        wanted = {*payloads, "@NOW()", "-ISO 4014", "=HYPERLINK(1)"}
        csv_texts = {v for row in csv_rows[1:] for v in row}
        safe_csv = all("'" + v in csv_texts for v in wanted) and "-12.5" in csv_texts
        passed = (
            pid_ok and safe_csv and len(csv_rows) == 5 and not formulas and wanted <= texts and numeric_mass
        )
        return strict_result(
            "export_safety",
            passed,
            {
                "pid_csv": line_rows,
                "bom_csv": csv_rows,
                "xlsx_texts": sorted(texts),
                "numeric_mass": numeric_mass,
                "formula_cells": formulas,
            },
        )

    async def task_stale_preview(self):
        await self.apply([{"op": "add_line", "start": [0, 0], "end": [10, 0], "tag": "line"}])
        ops = [{"op": "move", "target": "tag:line", "delta": [5, 0]}]
        preview = (await self.apply(ops, dry_run=True))["preview"]
        doc = self.read()
        doc.modelspace().query("LINE")[0].dxf.color = 5
        doc.saveas(self.workspace / self.path)
        handle = doc.modelspace().query("LINE")[0].dxf.handle
        applied = await self.apply(
            ops,
            expected_revision=preview["revision"],
            preconditions=preview["preconditions"],
            allow_error=True,
        )
        line = self.read().modelspace().query("LINE")[0]
        moved = tuple(line.dxf.start)[:2] != (0, 0)
        fresh = (await self.apply(ops, dry_run=True))["preview"]
        again = await self.apply(
            ops, expected_revision=fresh["revision"], preconditions=fresh["preconditions"]
        )
        final = list(self.read().entitydb[handle].dxf.start)[:2]
        passed = (
            not moved
            and not applied["ok"]
            and applied["error"]["code"] == "STALE_REVISION"
            and applied["error"]["entities"] == [handle]
            and again["ok"]
            and final == [5, 0]
        )
        return strict_result(
            "stale_preview", passed, {"refusal": applied, "moved_over_stale_preview": moved, "final": final}
        )

    async def task_revision_guarded_edit(self):
        await self.apply([{"op": "add_line", "start": [0, 0], "end": [10, 0], "tag": "line"}])
        ops = [{"op": "move", "target": "tag:line", "delta": [1, 0]}]
        revs = [(await self.call("read_drawing", path=self.path))["status"]["revision"]]
        for _ in range(2):
            result = await self.apply(ops, expected_revision=revs[-1])
            revs.append(result["revision"])
        before = digest(self.workspace / self.path)
        refused = await self.apply(ops, expected_revision=revs[0], allow_error=True)
        current = (await self.call("read_drawing", path=self.path))["status"]["revision"]
        error = refused.get("error", {})
        passed = (
            all(b == a + 1 for a, b in zip(revs, revs[1:]))
            and before == digest(self.workspace / self.path)
            and current == revs[-1]
            and error.get("code") == "STALE_REVISION"
            and error["details"]["expected_revision"] == revs[0]
            and error["details"]["actual_revision"] == revs[-1]
            and list(self.read().modelspace().query("LINE")[0].dxf.start)[:2] == [2, 0]
        )
        return strict_result("revision_guarded_edit", passed, {"revisions": revs, "refusal": refused})

    async def task_foreign_edit_verified(self):
        doc = ezdxf.new("R2010", units=4)
        msp = doc.modelspace()
        msp.add_lwpolyline([(0, 0), (100, 0), (100, 60), (0, 60)], close=True)
        holes = [("H1", 20, 20, 5), ("H2", 80, 20, 5), ("H3", 80, 40, 4.5), ("H4", 20, 40, 5)]
        for _, x, y, radius in holes:
            msp.add_circle((x, y), radius)
        msp.add_text("PLATE 100 x 60", height=2.5).set_placement((5, 52))
        doc.saveas(self.workspace / self.path)
        original = self.workspace / "foreign-original.dxf"
        doc.saveas(original)
        reqs = [{"id": name, "type": "hole", "center": [x, y], "diameter": 10} for name, x, y, _ in holes]
        reqs += [{"id": "N", "type": "hole_count", "count": 4}]
        before = await self.validate(requirements=reqs)
        selected = (await self.inspect(filter={"type": "CIRCLE"}))["entities"]
        pitch_handles = {tuple(e["center"][:2]): e["id"] for e in selected}
        pitch = await self.call(
            "apply_constraints",
            path=self.path,
            mode="check",
            constraints=[
                {
                    "id": "P",
                    "type": "distance",
                    "a": pitch_handles[(20, 20)] + "@center",
                    "b": pitch_handles[(80, 20)] + "@center",
                    "value": 60,
                }
            ],
        )
        fail = [r for r in before["requirements"] if r["status"] == "fail"]
        handle = fail[0]["entities"][0] if len(fail) == 1 else None
        if handle:
            fingerprint = next(e["fingerprint"] for e in selected if e["id"] == handle)
            await self.apply(
                [{"op": "modify", "target": handle, "set": {"radius": 5}}],
                preconditions={"entities": {handle: fingerprint}},
            )
        after = await self.validate(requirements=reqs)
        pitch_after = await self.call(
            "apply_constraints",
            path=self.path,
            mode="check",
            constraints=[
                {
                    "id": "P",
                    "type": "distance",
                    "a": pitch_handles[(20, 20)] + "@center",
                    "b": pitch_handles[(80, 20)] + "@center",
                    "value": 60,
                }
            ],
        )
        before_entities = {e.dxf.handle: e.dxfattribs() for e in ezdxf.readfile(original).modelspace()}
        after_entities = {e.dxf.handle: e.dxfattribs() for e in self.read().modelspace()}
        changed = [
            h for h in before_entities.keys() & after_entities if before_entities[h] != after_entities[h]
        ]
        passed = (
            [r["id"] for r in fail] == ["H3"]
            and all(r["status"] == "pass" for r in after["requirements"])
            and pitch["constraints"][0]["satisfied"]
            and pitch_after["constraints"][0]["satisfied"]
            and changed == [handle]
            and before_entities.keys() == after_entities.keys()
        )
        async with self.connect() as fresh:
            old_session, self.session = self.session, fresh
            try:
                reread = await self.validate(requirements=reqs)
                pitch_reopened = await self.call(
                    "apply_constraints",
                    path=self.path,
                    mode="check",
                    constraints=[
                        {
                            "id": "P",
                            "type": "distance",
                            "a": pitch_handles[(20, 20)] + "@center",
                            "b": pitch_handles[(80, 20)] + "@center",
                            "value": 60,
                        }
                    ],
                )
                passed &= (
                    all(r["status"] == "pass" for r in reread["requirements"])
                    and pitch_reopened["constraints"][0]["satisfied"]
                )
            finally:
                self.session = old_session
        return strict_result(
            "foreign_edit_verified",
            passed,
            {
                "before_failed": [r["id"] for r in fail],
                "after_failed": [r["id"] for r in after["requirements"] if r["status"] != "pass"],
                "pitch_before": pitch["constraints"],
                "pitch_after": pitch_after["constraints"],
                "pitch_reopened": pitch_reopened["constraints"],
                "changed": changed,
                "added": sorted(after_entities.keys() - before_entities),
                "removed": sorted(before_entities.keys() - after_entities),
                "fingerprint_guard": True,
            },
        )

    async def task_damaged_dxf_refusal(self):
        await self.apply([{"op": "add_line", "start": [0, 0], "end": [10, 0]}])
        before = digest(self.workspace / self.path)
        doc = ezdxf.new("R2010", units=4)
        for index in range(50):
            doc.modelspace().add_line((0, index), (10, index))
        intact = self.workspace / "intact.dxf"
        doc.saveas(intact)
        data = intact.read_bytes()
        (self.workspace / "truncated.dxf").write_bytes(data[: len(data) // 2])
        (self.workspace / "not-a-drawing.dxf").write_bytes(b"this is not a drawing\n" * 8)
        cases = {}
        for name in ["truncated.dxf", "not-a-drawing.dxf"]:
            response = await self.call("read_drawing", path=name, allow_error=True)
            err = response.get("error", {})
            cases[name] = {
                "refused": not response.get("ok"),
                "error": err,
                "named": "DXF" in json.dumps(err).upper(),
            }
        kept = before == digest(self.workspace / self.path) and len(self.read().modelspace()) == 1
        return strict_result(
            "damaged_dxf_refusal",
            kept and all(c["refused"] and c["named"] for c in cases.values()),
            {
                "cases": cases,
                "document_kept": kept,
                "identity": "file path + exact bytes; server is stateless",
            },
        )

    async def task_dimension_binding(self):
        # Identical four geometric/annotation traps to the pinned reference.
        doc = ezdxf.new("R2010", setup=True, units=4)
        msp = doc.modelspace()
        handles = {}
        for name, x, width, kwargs in [
            ("D1", 0, 42, {}),
            ("D2", 100, 40, {"text": "42"}),
            ("D3", 200, 21, {"override": {"dimlfac": 2}}),
        ]:
            msp.add_line((x, 0), (x, 20))
            msp.add_line((x + width, 0), (x + width, 20))
            dim = msp.add_linear_dim(
                base=(x, 30), p1=(x, 0), p2=(x + width, 0), dimstyle="Standard", **kwargs
            )
            dim.render()
            handles[name] = dim.dimension.dxf.handle
        msp.add_circle((350, 10), 20)
        dim = msp.add_diameter_dim(center=(350, 10), radius=20, angle=0, dimstyle="Standard")
        dim.render()
        handles["D4"] = dim.dimension.dxf.handle
        doc.saveas(self.workspace / self.path)
        expected = {"D1": 42, "D2": 40, "D3": 21, "D4": 40}
        report = await self.call(
            "verify_dimensions",
            path=self.path,
            requirements=[
                {"id": name, "handle": handles[name], "value": value} for name, value in expected.items()
            ],
        )
        rows = {r["id"]: r for r in report["rows"]}
        findings = {r["requirement"]: r["kind"] for r in report["findings"]}
        passed = (
            set(rows) == set(expected)
            and all(
                rows[n]["status"] == "pass" and abs(rows[n]["measured"] - v) < 1e-9
                for n, v in expected.items()
            )
            and findings == {"D2": "dimension_mismatch", "D3": "dimension_mismatch"}
            and rows["D2"]["displayed"] == rows["D3"]["displayed"] == 42
            and report["accepted"] is False
        )
        return strict_result("dimension_binding", passed, report)


def grade_pid(result):
    graph = result["graph"]
    rows = result["line_list"]
    numbers = {(r["from_tag"], r["to_tag"]): r["line_number"] for r in rows}
    authored = "100-P-1001-CS1"
    signal = [r for r in rows if r["line_class"] == "electric"]
    sequenced = [r for r in rows if r["line_class"] != "electric" and r["line_number"] != authored]
    return (
        len(graph["nodes"]) == 5
        and len(graph["edges"]) == len(rows) == 4
        and graph["stats"] == {"dangling": 0, "confidence_min": 1.0}
        and result["critique"] == []
        and result["instrument_index"] == [{"tag": "FIC-101", "connected_to": "FCV-101"}]
        and numbers.get(("P-101", "FCV-101")) == numbers.get(("FCV-101", "V-201")) == authored
        and len(signal) == len(sequenced) == 1
        and signal[0]["line_number"] is None
        and sequenced[0]["line_number"] == "3"
        and sequenced[0]["number_source"] == "xdata"
    )


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--upstream", type=Path, required=True)
    parser.add_argument(
        "--out", type=Path, required=True, help="New or empty directory (never overwrites evidence)"
    )
    parser.add_argument("--python", type=Path, default=Path(sys.executable))
    parser.add_argument("--task", action="append")
    parser.add_argument("--timeout", type=float, default=180)
    args = parser.parse_args()
    upstream = args.upstream.resolve()
    if git(upstream, "rev-parse", "HEAD") != PIN or git(
        upstream, "status", "--porcelain", "--untracked-files=no"
    ):
        parser.error(f"upstream must be a clean checkout of {PIN}")
    if args.out.exists() and any(args.out.iterdir()):
        parser.error("--out must be new or empty; use a new directory to retry")
    args.out.mkdir(parents=True, exist_ok=True)
    sys.path.insert(0, str(upstream))
    from benchmarks.run_competitors import run_tasks, _summary
    from benchmarks.tasks_v8 import TASKS_V8, TASKS_V2

    tasks = [t for t in TASKS_V8 if not args.task or t.task_id in args.task]
    if args.task and set(args.task) - {t.task_id for t in tasks}:
        parser.error("unknown task id")
    adapter = AgentCADAdapter(args.python.resolve())
    report = asyncio.run(run_tasks(adapter, tasks, args.out / "artifacts", timeout=args.timeout, matrix="v8"))
    # Upstream's manifest names its own repository. Explicitly identify the measured
    # product and retain the upstream manifest under evaluator instead of relabelling it.
    evaluator = report["manifest"]
    report["git_sha"] = adapter.product_sha
    report["manifest"] = {
        "server_sha": adapter.product_sha,
        "dirty": adapter.product_dirty,
        "source_sha256": adapter.source_hashes,
        "evaluator": evaluator,
        "adapter_sha256": digest(Path(__file__)),
        "dependencies": {name: importlib.metadata.version(name) for name in ["mcp", "ezdxf", "pypdf"]},
    }
    from benchmarks.adapters.base import TaskResult

    v2ids = {t.task_id for t in TASKS_V2}
    v2results = [TaskResult(**r) for r in report["results"] if r["task_id"] in v2ids]
    report["v2_summary"] = _summary(TASKS_V2, v2results)
    report["fixed_matrix_full_pass_percent"] = round(report["summary"]["passed"] / len(tasks) * 100, 2)
    report["evidence_boundary"] = (
        "Local independent adapter; not an upstream-certified leaderboard submission. Partial gates score 0."
    )
    for result in report["results"]:
        for artifact in result["artifacts"]:
            artifact.pop("path", None)
    write_json(args.out / "report.json", report)
    print(json.dumps({"summary": report["summary"], "v2_summary": report["v2_summary"]}, indent=2))


if __name__ == "__main__":
    main()
